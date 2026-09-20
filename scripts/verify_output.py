# -*- coding: utf-8 -*-
"""
RFI01 - 校验整理后的 Form Listing 输出表

用法:
    python verify_output.py "C:/Users/nomo/Downloads/20260920-处理过Form Listing.xlsx"

校验项:
  1) 结构：列数 21、表头在第 7 行、Discipline_Code 是否纯 BS
  2) 残留 AR / STR 单号
  3) 多日期分隔符是否为纯逗号（与源表 Response_Date 一致）
  4) 淡黄 7 列 / 淡粉 6 列的底色与边框
  5) Asite search 是否为文本格式（保留前导 0）
  6) ★ 签名反查：每段回复末尾署名(Best Regards, 人名)反查公司，与落位列比对
"""
import re, sys, collections, argparse
import openpyxl

HR = 7                       # 表头行
COL = {'asite': 4, 'title': 5, 'status': 6, 'ref': 8, 'bl': 9, 'floor': 10,
       'disc': 11, 'resp': 12, 'r_wtp': 13, 'r_wsp': 14, 'r_msc': 15,
       'created_by': 16, 'd_wtp': 17, 'd_wsp': 18, 'd_msc': 19, 'rd': 20}

COMP = {'WSP (Asia) Ltd': 'WSP', 'Wong Tung': 'WTP', 'MCS': 'MSC',
        'WSP': 'WSP', 'CSCE': 'CSHK'}
SIG = re.compile(
    r'(?:Best\s+Regards|Best\s+regards|Regards|Thanks\s*&\s*Regards|'
    r'Thanks\s+and\s+Regards|Thank\s+you|Thanks)\s*[,:]?\s*([A-Za-z][A-Za-z .\'-]{2,40})')
SEP_REPLY = '\n\n----------\n\n'


def build_name_map(ws):
    """从 Created_By 反推 人名 -> 公司代号"""
    pm = {}
    for r in range(HR + 1, ws.max_row + 1):
        toks = [p.strip() for p in str(ws.cell(row=r, column=COL['created_by']).value or '').split(',')]
        pend = None
        for t in toks:
            if t in COMP:
                if pend: pm[pend.lower()] = COMP[t]
                pend = None
            elif t and t != '---':
                pend = t
    return pm


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('path')
    a = ap.parse_args()
    ws = openpyxl.load_workbook(a.path).active
    n = ws.max_row - HR
    print('数据行 %d | 列数 %d' % (n, ws.max_column))

    head = [ws.cell(row=HR, column=i).value for i in range(1, ws.max_column + 1)]
    if len(head) != 21:
        print('!! 列数不是 21：', len(head))
    print('表头:', head)

    # 1) Discipline_Code
    print('\n[1] Discipline_Code:',
          dict(collections.Counter(ws.cell(row=r, column=COL['disc']).value
                                   for r in range(HR + 1, ws.max_row + 1))))

    # 2) 残留 AR / STR
    left = [ws.cell(row=r, column=COL['ref']).value for r in range(HR + 1, ws.max_row + 1)
            if re.search(r'/AR/|/STR/', str(ws.cell(row=r, column=COL['ref']).value or ''))]
    print('[2] 残留 AR/STR 单号:', len(left), left[:5])

    # 3) 日期分隔符
    slash = sum(1 for r in range(HR + 1, ws.max_row + 1) for c in (17, 18, 19)
                if '/' in str(ws.cell(row=r, column=c).value or ''))
    multi = sum(1 for r in range(HR + 1, ws.max_row + 1) for c in (17, 18, 19)
                if ',' in str(ws.cell(row=r, column=c).value or ''))
    print('[3] 多日期单元格 %d，残留斜杠 %d' % (multi, slash))

    # 4) 底色 / 边框
    Y, P = 'FFFFF2CC', 'FFFCE4EC'
    ok_y = ok_p = 0
    for i in (4, 13, 14, 15, 17, 18, 19):
        if ws.cell(row=HR + 1, column=i).fill.fgColor.rgb == Y: ok_y += 1
    for i in (5, 6, 8, 9, 10, 20):
        if ws.cell(row=HR + 1, column=i).fill.fgColor.rgb == P: ok_p += 1
    print('[4] 淡黄列 %d/7，淡粉列 %d/6，边框示例 %s'
          % (ok_y, ok_p, ws.cell(row=HR + 1, column=8).border.left.style))

    # 5) Asite search 文本格式
    print('[5] Asite search number_format =',
          repr(ws.cell(row=HR + 1, column=4).number_format),
          '| 样例', [ws.cell(row=HR + 1 + k, column=4).value for k in range(3)])

    # 6) 签名反查
    pm = build_name_map(ws)
    print('[6] 人名->公司映射 %d 条' % len(pm))
    agree = collections.Counter()
    for r in range(HR + 1, ws.max_row + 1):
        for c, code in ((13, 'WTP'), (14, 'WSP'), (15, 'MSC')):
            v = ws.cell(row=r, column=c).value
            if not v: continue
            for seg in str(v).split(SEP_REPLY):
                ms = list(SIG.finditer(seg))
                if not ms:
                    agree['无签名(不计)'] += 1
                    continue
                nm = ms[-1].group(1).strip().rstrip('.').lower()
                hit = pm.get(nm)
                if hit is None:
                    # 再试：在已知人名里找最长匹配
                    for k in sorted(pm, key=len, reverse=True):
                        if k in nm or nm in k:
                            hit = pm[k]; break
                if hit == code: agree['一致'] += 1
                elif hit is None: agree['人名未识别(不计)'] += 1
                else: agree['不一致'] += 1
    print('    ', dict(agree))
    tot = agree['一致'] + agree['不一致']
    if tot:
        print('    签名可判定段落一致率: %.2f%% (%d/%d)'
              % (100 * agree['一致'] / tot, agree['一致'], tot))

    # 统计
    print('\n--- 填值统计 ---')
    for k, nm in (('asite', 'Asite search'), ('r_wtp', 'Reply WTP'),
                  ('r_wsp', 'Reply WSP'), ('r_msc', 'Reply MSC'),
                  ('d_wtp', 'Date WTP'), ('d_wsp', 'Date WSP'), ('d_msc', 'Date MSC')):
        v = sum(1 for r in range(HR + 1, ws.max_row + 1)
                if ws.cell(row=r, column=COL[k]).value not in (None, ''))
        print('  %-14s %d' % (nm, v))
    print('  三项回复全空     %d' % sum(
        1 for r in range(HR + 1, ws.max_row + 1)
        if not any(ws.cell(row=r, column=c).value for c in (13, 14, 15))))


if __name__ == '__main__':
    main()
