# -*- coding: utf-8 -*-
"""
RFI01 - 整理 Asite 导出的 Form Listing.xlsx

用法:
    python build_bs_listing.py
    python build_bs_listing.py --src "C:/path/Form Listing.xlsx"
    python build_bs_listing.py --src "..." --dst "C:/path/20260920-处理过Form Listing.xlsx"
    python build_bs_listing.py --keep-ar-str          # 不剔除 AR/STR 单号
    python build_bs_listing.py --discipline BS --discipline FS

处理步骤:
  1) Discipline_Code 只保留指定专业的行(默认 BS)
  2) 剔除 Cont_Ref_No 学科段为 AR / STR 的行(可用 --keep-ar-str 关闭)
  3) 表单标题左侧新增 Asite search 列(标题开头数字, 文本格式 @, 保留前导 0)
  4) Response 右侧插入   Asite Reply (WTP) / (WSP) / (MSC)
  5) Created_By 右侧插入 Asite Reply Date (WTP) / (WSP) / (MSC)
  6) 按 Created_By 中顾问出现顺序(左->右) 与 Response_Date 日期(左->右) 按位配对
  7) Response 多段回复按分隔符 ' , ' 切分, 依 Created_By 顺序落位到对应顾问
  8) 按 Cont_Ref_No 自然排序
  9) 淡黄 7 列 + 淡粉 6 列: 底色 + 细格边框
"""
import re, os, sys, argparse, datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---------------- 常量 ----------------
# 公司代号：源表写法 -> 输出列名后缀。CSHK 是总包，不落列。
COMP = {'WSP (Asia) Ltd': 'WSP', 'Wong Tung': 'WTP', 'MCS': 'MSC',
        'WSP': 'WSP', 'CSCE': 'CSHK'}
COLS = ['WTP', 'WSP', 'MSC']                 # 落列的三家顾问
DELIM = ' , '                                # Response 多段回复的分隔符
RESP_SEP = '\n\n----------\n\n'              # 同一顾问多段回复的连接符
DATE_SEP = ','                               # 与源表 Response_Date 一致：纯逗号无空格
YELLOW = PatternFill('solid', fgColor='FFFFF2CC')
PINK = PatternFill('solid', fgColor='FFFCE4EC')
GRID = Border(*[Side(style='thin', color='FF808080')] * 4)

DROP_PREFIX = ('NAH-A/CS/RFI/AR/', 'NAH-A/CS/RFI/STR/')

# 列索引（源表 14 列，0-based）
I_TITLE, I_DISC, I_RESP, I_CB, I_RD = 3, 9, 10, 11, 12
I_REF = 6

HDR_ROW = 6          # 第 7 行是表头（0-based 索引 6）
DATA_START = 7       # 第 8 行起是数据


def parse_cb(cb):
    """Created_By: '人名, 公司,人名, 公司, …' -> [(人名, 公司代号), …]"""
    toks = [p.strip() for p in str(cb or '').split(',')]
    out, pend = [], None
    for t in toks:
        if t in COMP:
            out.append((pend, COMP[t])); pend = None
        elif t and t != '---':
            pend = t
    return out


def make_splitter(names):
    """构造 Response 切分器。names 是人名集合，用于给切点打分。"""
    NAME_RE = re.compile(r'^(' + '|'.join(re.escape(n) for n in
                         sorted(names, key=len, reverse=True)) + r')\b') if names else None
    HEAD_RE = re.compile(r'^(WTPL|WTP|WSP|MCS|CSHK|CS)\b', re.I)
    CLOSE_RE = re.compile(r'^(Close|Closing|Closed)\b', re.I)
    START_RE = re.compile(r'^(The|This|Please|As|RFI|Attachment|Subject|Note|We|I|'
                          r'Refer|According|Contract|Kindly|It)\b')

    def score_split(resp, pos):
        before = resp[:pos].rstrip()
        after = resp[pos + len(DELIM):].lstrip()
        s = 0
        if after.startswith('Dear'): s += 5
        if NAME_RE and NAME_RE.match(after): s += 5
        if HEAD_RE.match(after): s += 4
        if CLOSE_RE.match(after): s += 4
        if START_RE.match(after): s += 3
        if re.match(r'^[A-Z]', after): s += 1
        if before.endswith('.') or (NAME_RE and NAME_RE.search(before[-40:])): s += 3
        if after[:1] in '(-' or after[:1].islower(): s -= 5
        if before.endswith(',') or before.endswith(':'): s -= 3
        if re.match(r'^(i|ii|iii|iv|v|vi)\.', after) or re.match(r'^\d', after): s -= 4
        if re.match(r'^[a-z]', after): s -= 2
        return s

    def split_dp(resp, n):
        cand = [m.start() for m in re.finditer(re.escape(DELIM), resp)]
        if len(cand) < n - 1:
            return None
        sc = [score_split(resp, p) for p in cand]
        NEG = -10 ** 9
        dp = [[NEG] * n for _ in range(len(cand) + 1)]
        par = [[False] * n for _ in range(len(cand) + 1)]
        dp[0][0] = 0
        for j in range(1, len(cand) + 1):
            for k in range(0, min(j, n - 1) + 1):
                if dp[j - 1][k] > dp[j][k]:
                    dp[j][k] = dp[j - 1][k]; par[j][k] = False
                if k > 0 and dp[j - 1][k - 1] > NEG:
                    v = dp[j - 1][k - 1] + sc[j - 1]
                    if v > dp[j][k]:
                        dp[j][k] = v; par[j][k] = True
        if dp[len(cand)][n - 1] <= NEG:
            return None
        picks, j, k = [], len(cand), n - 1
        while j > 0:
            if par[j][k]:
                picks.append(cand[j - 1]); k -= 1
            j -= 1
        picks.reverse()
        segs, prev = [], 0
        for p in picks:
            segs.append(resp[prev:p].strip()); prev = p + len(DELIM)
        segs.append(resp[prev:].strip())
        return [s for s in segs if s]

    def split_resp(resp, n):
        """先按 ' , ' 直切；段数不符就用带评分的 DP 挑 n-1 个切点。"""
        if n <= 0:
            return []
        resp = str(resp or '')
        if not resp.strip() or resp.strip() == '---':
            return []
        s = [x.strip() for x in resp.split(DELIM) if x.strip()]
        if len(s) == n:
            return s
        d = split_dp(resp, n)
        return d if d else s

    return split_resp


def natkey(s):
    return [(0, int(p)) if p.isdigit() else (1, p.lower())
            for p in re.split(r'(\d+)', str(s or ''))]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default=os.path.join(
        os.path.expanduser('~'), 'Downloads', 'Form Listing.xlsx'),
        help='默认 ~/Downloads/Form Listing.xlsx')
    ap.add_argument('--dst', default=None,
                    help='默认 <下载目录>/<YYYYMMDD>-处理过Form Listing.xlsx')
    ap.add_argument('--sheet', default='Form Listing')
    ap.add_argument('--discipline', action='append', default=None,
                    help='保留的专业代码，可重复。默认 BS')
    ap.add_argument('--keep-ar-str', action='store_true',
                    help='保留 Cont_Ref_No 为 AR/STR 的行')
    ap.add_argument('--date', default=None, help='输出文件名日期，默认今天')
    a = ap.parse_args()

    disc = a.discipline or ['BS']
    day = a.date or datetime.date.today().strftime('%Y%m%d')
    if a.dst:
        DST = a.dst
    else:
        DST = os.path.join(os.path.dirname(a.src), '%s-处理过Form Listing.xlsx' % day)

    # ---------- 1. 读源数据 ----------
    wb = openpyxl.load_workbook(a.src, read_only=True)
    ws = wb[a.sheet]
    rows = list(ws.iter_rows(values_only=True))
    HEAD = list(rows[HDR_ROW])
    ALL = [list(r) for r in rows[DATA_START:]]
    DATA = [r for r in ALL if str(r[I_DISC]).strip() in disc]
    print('源数据行 %d, 其中 %s 行 %d' % (len(ALL), '+'.join(disc), len(DATA)))

    # ---------- 2. 剔除 AR / STR 单号 ----------
    if not a.keep_ar_str:
        before = len(DATA)
        DATA = [r for r in DATA
                if not any(str(r[I_REF] or '').startswith(p) for p in DROP_PREFIX)]
        print('剔除 AR/STR 后剩余 %d 行（删 %d）' % (len(DATA), before - len(DATA)))

    # ---------- 3. 解析 Created_By ----------
    ENT = [parse_cb(r[I_CB]) for r in DATA]
    bad = [(r[I_REF], r[I_CB]) for r, e in zip(DATA, ENT)
           if sum(1 for x in str(r[I_CB]).split(',') if x.strip() in COMP) != len(e)]
    print('Created_By 解析校验:', ('!! 异常 %d 条' % len(bad)) if bad else '通过')
    for b in bad[:10]:
        print('   ', b)

    pm = {}
    for e in ENT:
        for n, c in e:
            if n:
                pm[n] = c
    print('人名->公司映射 %d 条' % len(pm))

    split_resp = make_splitter(set(pm))

    # ---------- 4. 汇总每行 ----------
    out_rows = []
    for r, ent in zip(DATA, ENT):
        n = len(ent)
        segs = split_resp(r[I_RESP], n)
        dates = [d.strip() for d in str(r[I_RD] or '').split(',') if d.strip()]
        reply = {c: [] for c in COLS}
        date = {c: [] for c in COLS}
        for i in range(n):
            code = ent[i][1]
            if code not in reply:
                continue                       # CSHK 等总包，不落列
            if i < len(segs): reply[code].append(segs[i])
            if i < len(dates): date[code].append(dates[i])
        m = re.match(r'^(\d+)', str(r[I_TITLE] or '').strip())
        asite = m.group(1) if m else None
        new = list(r[:3])                                        # A-C
        new += [asite]                                           # D  Asite search
        new += [r[3]]                                            # E  表单标题
        new += list(r[4:10])                                     # F-K
        new += [r[10]]                                           # L  Response
        new += [RESP_SEP.join(reply[c]) or None for c in COLS]   # M,N,O
        new += [r[11]]                                           # P  Created_By
        new += [DATE_SEP.join(date[c]) or None for c in COLS]    # Q,R,S
        new += [r[12], r[13]]                                    # T,U
        out_rows.append(new)

    # ---------- 5. 按 Cont_Ref_No 自然排序 ----------
    def sortkey(x):
        ref = str(x[7] or '')                       # 插入 Asite search 后索引变 7！
        broken = 0 if re.match(r'^NA[FH]', ref, re.I) else 1
        return (broken, natkey(ref))
    out_rows.sort(key=sortkey)
    print('排序后首条:', out_rows[0][7], '| 末条:', out_rows[-1][7])

    # ---------- 6. 写新表 ----------
    nwb = openpyxl.Workbook()
    nws = nwb.active
    nws.title = a.sheet
    for i, v in enumerate(rows[2]):
        if v is not None: nws.cell(row=3, column=i + 1, value=v)
    for i, v in enumerate(rows[3]):
        if v is not None: nws.cell(row=4, column=i + 1, value=v)

    NEWHEAD = (HEAD[:3] + ['Asite search'] + [HEAD[3]] + HEAD[4:10] + ['Response']
               + ['Asite Reply (WTP)', 'Asite Reply (WSP)', 'Asite Reply (MSC)']
               + ['Created_By']
               + ['Asite Reply Date (WTP)', 'Asite Reply Date (WSP)',
                  'Asite Reply Date (MSC)']
               + ['Response_Date', 'Response_Flag'])
    assert len(NEWHEAD) == 21, len(NEWHEAD)
    HR = HDR_ROW + 1                                  # 1-based 表头行号 = 7
    for i, h in enumerate(NEWHEAD, start=1):
        c = nws.cell(row=HR, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical='top')

    for j, row in enumerate(out_rows):
        for i, v in enumerate(row, start=1):
            c = nws.cell(row=HR + 1 + j, column=i, value=v)
            c.alignment = Alignment(vertical='top', wrap_text=(i in (12, 13, 14, 15)))

    HILITE_YELLOW = [4, 13, 14, 15, 17, 18, 19]
    HILITE_PINK = [5, 6, 8, 9, 10, 20]
    for cols, fill in ((HILITE_YELLOW, YELLOW), (HILITE_PINK, PINK)):
        for i in cols:
            for r in range(HR, HR + len(out_rows) + 1):
                c = nws.cell(row=r, column=i)
                c.fill = fill
                c.border = GRID
    for r in range(HR + 1, HR + len(out_rows) + 1):   # Asite search 强制文本格式
        nws.cell(row=r, column=4).number_format = '@'

    wid = {1: 23.4, 2: 23.4, 3: 13, 4: 12, 5: 46.9, 6: 15.6, 7: 13, 8: 46.9, 9: 46.9,
           10: 46.9, 11: 13, 12: 60, 13: 60, 14: 60, 15: 46.9,
           16: 46.9, 17: 20, 18: 20, 19: 20, 20: 46.9, 21: 13}
    for k, w in wid.items():
        nws.column_dimensions[get_column_letter(k)].width = w
    nws.freeze_panes = 'A%d' % (HR + 1)

    try:
        nwb.save(DST)
    except PermissionError:
        print('\n!! 无法写入：%s 正被 Excel/WPS 打开。请关闭后重跑。' % DST)
        print('   （目录下会有 ~$<文件名> 锁文件，确认它消失即可）')
        sys.exit(1)
    print('已保存 ->', DST, '| 数据行', len(out_rows), '| 列数', nws.max_column)


if __name__ == '__main__':
    main()
